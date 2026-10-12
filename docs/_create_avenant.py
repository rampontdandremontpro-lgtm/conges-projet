# -*- coding: utf-8 -*-
"""Création de l'avenant fonctionnel au cahier des charges — 12/10/2026."""
import docx
from docx.shared import Pt

OUT = 'docs/avenant-cahier-des-charges-2026-10-12.docx'

d = docx.Document()

def H(text, level=1):
    d.add_heading(text, level=level)

def P(text):
    d.add_paragraph(text)

def bullet(text):
    d.add_paragraph(text, style='List Bullet')

H('Avenant fonctionnel au cahier des charges', 0)
P('G Congés & Absences — Évolution fonctionnelle validée le 12/10/2026')

H('1. Objet', 1)
P('Le présent avenant remplace, à compter du 12/10/2026, les dispositions contradictoires du cahier des charges initial concernant la déclaration des absences administratives et le calcul du prorata (BAL-009). Le document historique d’origine reste conservé comme référence.')

H('2. Déclarations d’absence — exclusivité RH', 1)
P('Nouvelle règle : seule la RH peut créer, soumettre, modifier et annuler une déclaration administrative d’absence.')
bullet('Le collaborateur ne peut plus créer, soumettre ni annuler une déclaration d’absence, même pour lui-même (HTTP 403, aucune mutation).')
bullet('Le collaborateur conserve : la consultation de ses propres absences, la transmission de justificatifs autorisés et le remplacement d’un justificatif rejeté.')
bullet('Le Responsable de service, le Directeur (module RH) et l’Administrateur ne créent pas de déclaration d’absence administrative.')
bullet('Les demandes de congés restent inchangées et accessibles selon les droits existants.')
bullet('Les fonctions spécifiques d’indisponibilité du Directeur restent fonctionnelles et distinctes des déclarations administratives réservées à la RH.')

H('3. Prorata d’acquisition (BAL-009)', 1)
P('Règle fonctionnelle définitivement arbitrée : prorata mensuel = taux mensuel configuré × nombre de jours calendaires de présence inclusifs dans le mois ÷ nombre réel de jours calendaires de ce mois.')
bullet('Jour d’arrivée inclus ; jour de départ inclus.')
bullet('Arrondi mathématique à deux décimales ; crédit automatique.')
bullet('Date de sortie enregistrée sur le collaborateur.')
bullet('Correction RH toujours possible par mouvement séparé et tracé ; l’historique des acquisitions est conservé.')
P('Cas de référence (taux 2,50 jours) : mois complet 2,50 ; arrivée le 16 juillet 1,29 ; arrivée le 31 juillet 0,08 ; départ le 15 juillet 1,21 ; départ le 31 juillet 2,50 ; arrivée le 16 et départ le 20 juillet 0,40.')
P('Cet arbitrage est une règle applicative GMES et ne constitue pas une obligation légale générale.')

H('4. Impacts sur les droits et parcours', 1)
bullet('Frontend : suppression du bouton « Déclarer une absence » du tableau de bord collaborateur et retrait de la route /app/new-absence.')
bullet('Backend : routes POST /absence-declarations, POST /:id/submit, POST /:id/cancel, PATCH /:id et DELETE /:id réservées au rôle RH (HTTP 403 pour les autres rôles).')
bullet('Lecture et dépôt de justificatifs collaborateurs inchangés.')
bullet('Absences existantes et historiques non modifiés rétroactivement.')

H('5. Preuves techniques associées', 1)
bullet('retest-absence-rh-only.json (28/28 API)')
bullet('retest-absence-rh-only-ui.json (8/8 Playwright)')
bullet('micro-validation-absence-rh-only.json (11/11)')
bullet('retest-abs-rh-only-officiel.json (17/17 ABS)')
bullet('retest-bal009-prorata.json et retest-bal009-ui.json (BAL-009)')

H('6. Validation', 1)
P('Date : 12/10/2026. Cet avenant s’applique à l’ensemble du périmètre de la recette et de la mise en production de G Congés & Absences.')

d.save(OUT)
print('Avenant créé :', OUT)
