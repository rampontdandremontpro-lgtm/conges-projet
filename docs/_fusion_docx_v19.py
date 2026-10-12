# -*- coding: utf-8 -*-
"""Fusion documentaire finale — cahier de recette v1.9 (12/10/2026)."""
import docx

DOCX = 'docs/CAHIER DE RECETTE - DÉMATÉRIALISATION DES CONGÉS DES ABSENCES.docx'
DATE = '12/10/2026'
PROOF = 'retest-abs-rh-only-officiel.json'

# Contenu du nouveau référentiel ABS (colonnes : Scénario, Préconditions/procédure, Résultat attendu, Preuve, Résultat observé)
ABS = {
 'ABS-001': ('Collaborateur ne peut pas créer une déclaration (même pour lui-même)',
   'Préconditions : compte Collaborateur actif, type d’absence actif. Étapes : 1. Le collaborateur tente de créer une déclaration d’absence pour lui-même (POST /absence-declarations). 2. Vérifier le code HTTP et l’absence de mutation.',
   'HTTP 403 ; aucune déclaration créée ; aucune mutation en base.',
   PROOF,
   'POST /absence-declarations (soi-même) HTTP=403 ; aucune création, aucune mutation BDD. Statut : Conforme'),
 'ABS-002': ('Collaborateur ne crée pas une absence pour un collègue',
   'Préconditions : comptes Collaborateur col-b et col-c actifs. Étapes : 1. col-b tente POST /absence-declarations avec employeeId=col-c. 2. Vérifier le refus et l’absence de mutation.',
   'HTTP 403 ; aucune mutation.',
   PROOF,
   'POST /absence-declarations (employeeId=COL-C) HTTP=403 ; aucune mutation BDD. Statut : Conforme'),
 'ABS-003': ('Tous les types de déclaration d’absence sont réservés à la RH (création)',
   'Préconditions : type normal et type à justificatif actifs. Étapes : 1. Le collaborateur tente de créer avec un type normal puis un type à justificatif. 2. Vérifier le refus.',
   'HTTP 403 quel que soit le type ; aucune mutation.',
   PROOF,
   'POST type normal HTTP=403 ; type justificatif HTTP=403 ; aucune mutation BDD. Statut : Conforme'),
 'ABS-004': ('Collaborateur ne peut plus soumettre un brouillon (même ancien, lui appartenant)',
   'Préconditions : brouillon BROUILLON préparé par la RH pour le collaborateur. Étapes : 1. Le collaborateur tente POST /absence-declarations/:id/submit. 2. Vérifier le refus et le statut inchangé.',
   'HTTP 403 ; statut BROUILLON inchangé.',
   PROOF,
   'POST /absence-declarations/:id/submit HTTP=403 ; statut BROUILLON inchangé. Statut : Conforme'),
 'ABS-005': ('Autre collaborateur ne consulte pas cette absence',
   'Préconditions : absence créée par la RH pour col-b. Étapes : 1. col-c tente GET /absence-declarations/:id. 2. col-b (propriétaire) consulte.',
   'col-b 200 ; col-c 403/404.',
   PROOF,
   'GET /absence-declarations/:id : propriétaire HTTP=200, autre collaborateur HTTP=403. Statut : Conforme'),
 'ABS-006': ('Collaborateur ajoute un justificatif PDF valide',
   'Préconditions : absence JUSTIFICATIF_EN_ATTENTE créée par la RH. Étapes : 1. Le collaborateur dépose un justificatif PDF. 2. Vérifier le statut.',
   'HTTP 201 ; absence → A_VERIFIER_PAR_RH ; document EN_ATTENTE.',
   PROOF,
   'POST /documents/absence/:id HTTP=201 ; absence A_VERIFIER_PAR_RH ; document EN_ATTENTE. Statut : Conforme'),
 'ABS-007': ('Autre collaborateur ne consulte pas le justificatif',
   'Préconditions : justificatif déposé (ABS-006). Étapes : 1. Un autre collaborateur tente de lister/télécharger le justificatif.',
   'HTTP 403/404 ; aucun PDF divulgué.',
   PROOF,
   'liste API HTTP=403 ; download API HTTP=403 ; aucun PDF divulgué. Statut : Conforme'),
 'ABS-008': ('RH consulte les justificatifs à contrôler',
   'Préconditions : document EN_ATTENTE. Étapes : 1. La RH consulte les documents à contrôler.',
   'HTTP 200 ; document EN_ATTENTE présent.',
   PROOF,
   'GET /documents/management?status=EN_ATTENTE HTTP=200 ; document EN_ATTENTE présent. Statut : Conforme'),
 'ABS-009': ('Directeur ne télécharge pas le justificatif',
   'Préconditions : document EN_ATTENTE. Étapes : 1. Le Directeur tente le téléchargement.',
   'HTTP 403/404 ; aucun PDF divulgué.',
   PROOF,
   'GET /documents/:id/download HTTP=403 ; aucun PDF divulgué. Statut : Conforme'),
 'ABS-010': ('RH télécharge le justificatif',
   'Préconditions : document EN_ATTENTE. Étapes : 1. La RH télécharge le justificatif.',
   'HTTP 200 ; PDF conforme (signature + hash).',
   PROOF,
   'GET /documents/:id/download HTTP=200 ; Content-Type=application/pdf ; hash SHA-256 identique. Statut : Conforme'),
 'ABS-011': ('RH rejette le justificatif avec un motif',
   'Préconditions : document EN_ATTENTE. Étapes : 1. La RH rejette avec motif.',
   'HTTP 200 ; document REJETE ; absence → JUSTIFICATIF_EN_ATTENTE.',
   PROOF,
   'POST /documents/:id/reject HTTP=200 ; document REJETE ; absence JUSTIFICATIF_EN_ATTENTE. Statut : Conforme'),
 'ABS-012': ('Collaborateur remplace le justificatif rejeté',
   'Préconditions : justificatif REJETE. Étapes : 1. Le collaborateur transmet un document de remplacement.',
   'HTTP 201 ; ancien ARCHIVE ; nouveau EN_ATTENTE ; absence → A_VERIFIER_PAR_RH.',
   PROOF,
   'POST /documents/absence/:id HTTP=201 ; ancien ARCHIVE ; nouveau EN_ATTENTE ; absence A_VERIFIER_PAR_RH. Statut : Conforme'),
 'ABS-013': ('L’ancien justificatif archivé n’est plus traitable',
   'Préconditions : ancien document ARCHIVE. Étapes : 1. Tenter accept/reject sur l’ancien document.',
   'HTTP 404/400 ; non traitable.',
   PROOF,
   'accept ancien HTTP=404 ; reject ancien HTTP=404. Statut : Conforme'),
 'ABS-014': ('RH accepte le justificatif de remplacement',
   'Préconditions : nouveau document EN_ATTENTE. Étapes : 1. La RH accepte le justificatif.',
   'HTTP 200 ACCEPTE ; 2e acceptation refusée.',
   PROOF,
   'POST /documents/:id/accept HTTP=200 ACCEPTE ; 2e acceptation HTTP=400. Statut : Conforme'),
 'ABS-015': ('RH enregistre définitivement l’absence',
   'Préconditions : justificatif accepté, absence A_VERIFIER_PAR_RH. Étapes : 1. La RH enregistre définitivement.',
   'HTTP 200 ; absence → ENREGISTREE.',
   PROOF,
   'POST /absence-declarations/:id/register HTTP=200 ; statut ENREGISTREE. Statut : Conforme'),
 'ABS-016': ('RH crée une absence pour un collaborateur (interface réelle)',
   'Préconditions : compte RH actif, collaborateur actif, type d’absence actif. Étapes : 1. UI RH : Déclarer congé/absence → sélection collaborateur + type → date → Valider.',
   'Création 201 ; statut ENREGISTREE ; visible par le collaborateur, invisible pour autrui.',
   'Playwright + API + SQL — ' + PROOF,
   'UI RH création HTTP=201 ; statut ENREGISTREE ; visible par le collaborateur, invisible pour autrui. Statut : Conforme'),
 'ABS-017': ('Collaborateur ne peut plus annuler son brouillon d’absence',
   'Préconditions : brouillon BROUILLON préparé par la RH. Étapes : 1. Le collaborateur tente POST /absence-declarations/:id/cancel.',
   'HTTP 403 ; statut BROUILLON inchangé.',
   PROOF,
   'POST /absence-declarations/:id/cancel HTTP=403 ; statut BROUILLON inchangé. Statut : Conforme'),
}

BAL009_EXPECTED = 'Prorata calendaire : taux mensuel × jours calendaires de présence (arrivée et départ inclus) ÷ jours calendaires du mois ; arrondi mathématique à 2 décimales ; crédit automatique ; date de sortie enregistrée ; correction RH tracée.'
BAL009_PROOF = 'retest-bal009-prorata.json + retest-bal009-ui.json'
BAL009_RESULT = 'Cas de référence (taux 2,50) : complet 2,50 ; arrivée 16/07 1,29 ; arrivée 31/07 0,08 ; départ 15/07 1,21 ; départ 31/07 2,50 ; arrivée 16/07+départ 20/07 0,40. Idempotence et correction RH vérifiées. Statut : Conforme'

def set_cell(cell, text):
    # Remplace le texte en conservant la première structure de paragraphe.
    if cell.paragraphs:
        p = cell.paragraphs[0]
        for r in list(p.runs[1:]):
            r._element.getparent().remove(r._element)
        if p.runs:
            p.runs[0].text = text
        else:
            p.add_run(text)
        for extra in list(cell.paragraphs[1:]):
            extra._element.getparent().remove(extra._element)
    else:
        cell.text = text

def set_para(p, text):
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

def find_row(table, cell0_text):
    for r in table.rows:
        if r.cells[0].text.strip() == cell0_text:
            return r
    return None

d = docx.Document(DOCX)

# 1) Table 1 — Informations du document
info = d.tables[1]
for r in info.rows:
    k = r.cells[0].text.strip()
    if k == 'Version': set_cell(r.cells[1], '1.9')
    elif k == 'Date': set_cell(r.cells[1], DATE)
    elif k == 'Statut du document': set_cell(r.cells[1], 'Recette acceptée')

# 2) Table 2 — Historique des versions : ajouter v1.9
hist = d.tables[2]
new = hist.add_row()
vals = ['1.9', DATE, 'Daphné RAMPONT-D’ANDREMONT',
        'Évolution fonctionnelle RH-only (12/10/2026) + fusion BAL-009 (prorata arbitré).',
        'Recette acceptée — 420/420 scénarios conformes.']
for i, v in enumerate(vals):
    set_cell(new.cells[i], v)

# 3) Table 11 — 17 scénarios ABS
tabs = d.tables[11]
for r in tabs.rows:
    idc = r.cells[0].text.strip()
    if idc in ABS:
        scen, proc, exp, proof, result = ABS[idc]
        set_cell(r.cells[2], scen)
        set_cell(r.cells[4], proc)
        set_cell(r.cells[5], exp)
        set_cell(r.cells[6], proof)
        set_cell(r.cells[7], 'Résultat obtenu : ' + result)

# 4) Table 10 — BAL-009
tbal = d.tables[10]
row = find_row(tbal, 'BAL-009')
if row is not None:
    set_cell(row.cells[5], BAL009_EXPECTED)
    set_cell(row.cells[6], BAL009_PROOF)
    set_cell(row.cells[7], 'Résultat obtenu : ' + BAL009_RESULT)

# 5) Table 28 — Synthèse décision
syn = d.tables[28]
for r in syn.rows:
    k = r.cells[0].text.strip()
    if k == 'Scénarios conformes': set_cell(r.cells[1], '420')
    elif k == 'Scénarios bloqués': set_cell(r.cells[1], '0')
    elif k == 'Taux de réussite': set_cell(r.cells[1], '100 %')

# 6) Décision — paragraphes
def replace_in_para(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)

for i, p in enumerate(d.paragraphs):
    t = p.text
    if t.strip() == '☐ Recette validée sans réserve':
        replace_in_para(p, '☐', '☒')
    elif t.strip().startswith('☒ Recette refusée'):
        replace_in_para(p, '☒', '☐')
    elif t.strip() == 'Réserves ou commentaires généraux :':
        pass
    elif t.startswith('Recette initiale exécutée du'):
        set_para(p, 'Recette exécutée du 22/09/2026 au 12/10/2026 ; LOT 1 à LOT 6 ; évolution fonctionnelle RH-only validée le 12/10/2026 ; BAL-009 arbitré et validé le 12/10/2026 ; 420/420 scénarios conformes ; aucune anomalie ouverte.')

d.save(DOCX)
print('DOCX v1.9 enregistré :', DOCX)
